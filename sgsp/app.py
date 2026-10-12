"""
SGSP – State Government Salary Package awareness voice agent (outbound calls to employees)
==========================================================================================
Run:
    uvicorn sgsp.app:app --host 0.0.0.0 --port 8004 --reload

Endpoints (admin = needs header X-Admin-Token when SGSP_ADMIN_TOKEN is set):
    POST /sgsp/campaigns                 admin — upload an HRMS .xlsx / .csv (phone, name, …) → campaign
    GET  /sgsp/campaigns                 admin — all campaigns with live stats
    GET  /sgsp/campaigns/{id}            admin — one campaign
    POST /sgsp/campaigns/{id}/pause      admin — stop dialing (calls in progress finish)
    POST /sgsp/campaigns/{id}/resume     admin
    GET  /sgsp/campaigns/{id}/report.csv admin — one row per employee
    POST /sgsp/call                      admin — {"phone", "name", "gender"?} single call (?dial=false: no Exotel)
    POST /sgsp/dnc                       admin — {"phone"} add to the do-not-call list
    POST /sgsp/call-status               Exotel status callback (busy / no-answer → retry)
    WS   /sgsp/stream                    Exotel App Studio "Stream" applet connects here
    GET  /sgsp/dashboard                 campaign dashboard
    GET  /sgsp/web                       browser test page (talk to the agent, no phone)
    WS   /sgsp/web/stream                browser voice session (?name=&gender=)
"""

import asyncio
import csv
import io
import json
import logging
import os
import re
import sys
import time
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile, WebSocket
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel

from . import store
from .audio_utils import cleanup_call_logging, setup_call_logging
from .config import (
    ADMIN_TOKEN,
    CALL_RATE_LIMIT_MAX,
    CALL_RATE_LIMIT_WINDOW_S,
    CHANNEL_LIMIT,
    EXOTEL_CALLER_ID,
    EXOTEL_FLOW_URL,
    EXOTEL_STATUS_CALLBACK_URL,
    GEMINI_MODEL,
    LOGS_DIR,
    PREWARM,
    PREWARM_TTL_S,
    SMS_ENABLED,
    WEBHOOK_URL,
)
from .dispatcher import dispatcher_loop, within_calling_hours
from .live import LiveConnection
from .session import TOOLS, ExotelSession, WebSession

# Windows consoles default to cp1252, which can't print the log separators / Kannada.
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(message)s", datefmt="%H:%M:%S", stream=sys.stdout)
log = logging.getLogger("sgsp")

IST = timezone(timedelta(hours=5, minutes=30))
STATIC_DIR = Path(__file__).parent / "static"
# Exotel statuses meaning the call will never be answered.
UNANSWERED_STATUSES = {"busy", "no-answer", "failed", "canceled"}

# phone → {"call_id", "employee", "prewarm", "call_sid"} from dialing until Exotel streams in.
pending_calls: dict[str, dict] = {}
_call_attempts: dict[str, deque] = defaultdict(deque)


@asynccontextmanager
async def lifespan(_app: FastAPI):
    store.init()
    store.recover_after_restart()
    task = None
    if os.environ.get("SGSP_DISPATCHER", "true").lower() not in ("0", "false", "no", "off"):
        task = asyncio.create_task(dispatcher_loop(dial_employee))
    yield
    if task:
        task.cancel()


app = FastAPI(title="SGSP Salary Package Voice Agent", lifespan=lifespan)


def require_admin(x_admin_token: str | None = Header(default=None)):
    if ADMIN_TOKEN and x_admin_token != ADMIN_TOKEN:
        raise HTTPException(status_code=401, detail="Missing or wrong X-Admin-Token")


# ── Helpers ────────────────────────────────────────────────────────────────────
_MOBILE_RE = re.compile(r"^[6-9]\d{9}$")


def normalize_phone(raw) -> str | None:
    """Return +91XXXXXXXXXX for a valid Indian mobile, else None. Excel numbers arrive as floats."""
    if isinstance(raw, float) and raw.is_integer():
        raw = str(int(raw))
    s = re.sub(r"[\s\-()]", "", str(raw or ""))
    if s.startswith("+91"):
        s = s[3:]
    elif s.startswith("91") and len(s) == 12:
        s = s[2:]
    elif s.startswith("0") and len(s) == 11:
        s = s[1:]
    return f"+91{s}" if _MOBILE_RE.match(s) else None


def _rate_limited(phone: str) -> bool:
    now = time.time()
    attempts = _call_attempts[phone]
    while attempts and now - attempts[0] > CALL_RATE_LIMIT_WINDOW_S:
        attempts.popleft()
    if len(attempts) >= CALL_RATE_LIMIT_MAX:
        return True
    attempts.append(now)
    return False


def place_exotel_call(to_number: str) -> dict:
    """Outbound call via Exotel Connect API; the SGSP App Studio flow streams to /sgsp/stream."""
    sid = os.environ.get("EXOTEL_SID")
    api_key = os.environ.get("EXOTEL_API_KEY")
    api_token = os.environ.get("EXOTEL_API_TOKEN")
    domain = os.environ.get("EXOTEL_API_DOMAIN", "api.in.exotel.com")

    missing = [k for k, v in {
        "EXOTEL_SID": sid, "EXOTEL_API_KEY": api_key, "EXOTEL_API_TOKEN": api_token,
        "SGSP_EXOTEL_CALLER_ID/EXOTEL_CALLER_ID": EXOTEL_CALLER_ID, "SGSP_EXOTEL_FLOW_URL": EXOTEL_FLOW_URL,
    }.items() if not v]
    if missing:
        raise ValueError(f"Missing Exotel env vars: {', '.join(missing)}")

    data = {"From": to_number, "CallerId": EXOTEL_CALLER_ID, "Url": EXOTEL_FLOW_URL, "CallType": "trans"}
    if EXOTEL_STATUS_CALLBACK_URL:
        data["StatusCallback"] = EXOTEL_STATUS_CALLBACK_URL
    resp = requests.post(f"https://{domain}/v1/Accounts/{sid}/Calls/connect.json",
                         auth=(api_key, api_token), data=data, timeout=15)
    resp.raise_for_status()
    call = resp.json().get("Call", {})
    return {"call_sid": call.get("Sid", ""), "status": call.get("Status", "")}


def _new_call_dir(label: str) -> tuple[str, Path]:
    run_ts = datetime.now(IST).strftime("%Y%m%d_%H%M%S")
    call_dir = LOGS_DIR / f"{run_ts}_{label}"
    call_dir.mkdir(parents=True, exist_ok=True)
    return run_ts, call_dir


def _employee(row: dict) -> dict:
    return {k: row.get(k) for k in ("name", "gender", "employee_id")}


async def _post_webhook(payload: dict):
    if not WEBHOOK_URL:
        return
    try:
        await asyncio.get_running_loop().run_in_executor(
            None, lambda: requests.post(WEBHOOK_URL, json=payload, timeout=10).raise_for_status())
    except Exception as e:
        log.error(f"[WEBHOOK] Delivery failed for call {payload.get('callId')}: {e}")


# ── Dialing (used by the dispatcher and POST /sgsp/call) ───────────────────────

async def _prewarm(phone: str, employee: dict) -> LiveConnection | None:
    try:
        conn = await LiveConnection(employee, "phone", TOOLS, log).open()
        conn.start_greeting()
        return conn
    except Exception as e:
        log.warning(f"[LIVE]  Pre-warm failed for {phone} — will connect when the call is answered: {e}")
        return None


async def _discard_prewarm(info: dict | None):
    task = (info or {}).get("prewarm")
    if not task:
        return
    try:
        conn = await asyncio.wait_for(task, timeout=15)
    except Exception:
        return
    if conn:
        await conn.close()


async def _expire_prewarm(phone: str, task: asyncio.Task):
    await asyncio.sleep(PREWARM_TTL_S)
    info = pending_calls.get(phone)
    if info and info.get("prewarm") is task:
        log.info(f"[CALL]  {phone} never streamed in within {PREWARM_TTL_S}s — dropping its pre-warmed session")
        await _discard_prewarm(pending_calls.pop(phone, None))


async def dial_employee(row: dict, dial: bool = True) -> dict:
    """Register the pending call (with a pre-warmed Gemini session) and ask Exotel to dial."""
    phone, employee = row["phone"], _employee(row)
    await _discard_prewarm(pending_calls.pop(phone, None))
    info = {"call_id": row["id"], "employee": employee, "requested_at": time.time()}
    if PREWARM:
        info["prewarm"] = asyncio.create_task(_prewarm(phone, employee))
        asyncio.create_task(_expire_prewarm(phone, info["prewarm"]))
    pending_calls[phone] = info
    if not dial:
        return {"status": "queued", "message": "Dial skipped (local mode)."}
    try:
        # requests is blocking — keep it off the event loop so the pre-warm isn't stalled.
        result = await asyncio.get_running_loop().run_in_executor(None, place_exotel_call, phone)
    except Exception as e:
        if isinstance(e, requests.HTTPError) and e.response is not None:
            body = f"Exotel HTTP {e.response.status_code}: {e.response.text.strip()[:500]}"
        else:
            body = f"{type(e).__name__}: {e}"
        log.error(f"[CALL]  Could not dial call {row['id']} ({phone}): {body}")
        await _discard_prewarm(pending_calls.pop(phone, None))
        store.dial_failed(row["id"], "failed")
        return {"status": "error", "detail": body}
    info["call_sid"] = result["call_sid"]
    store.set_call_sid(row["id"], result["call_sid"])
    log.info(f"[CALL]  Exotel call placed — call {row['id']} sid={result['call_sid']} status={result['status']}")
    return {"status": "dialing", "call_sid": result["call_sid"]}


# ── Endpoints ──────────────────────────────────────────────────────────────────

@app.get("/sgsp/health")
async def health():
    return {"status": "ok", "model": GEMINI_MODEL, "active_calls": store.active_call_count(),
            "channel_limit": CHANNEL_LIMIT, "calling_hours_open": within_calling_hours(),
            "sms_enabled": SMS_ENABLED}


class CallRequest(BaseModel):
    phone: str
    name: str
    gender: str | None = None
    employee_id: str | None = None


@app.post("/sgsp/call", dependencies=[Depends(require_admin)])
async def request_call(req: CallRequest, dial: bool = True):
    """Call one employee now (testing / manual follow-up). Ignores calling hours, respects DNC."""
    phone = normalize_phone(req.phone)
    if not phone:
        raise HTTPException(status_code=422, detail="Invalid Indian mobile number")
    if store.is_dnc(phone):
        raise HTTPException(status_code=409, detail="This number is on the do-not-call list")
    if _rate_limited(phone):
        raise HTTPException(status_code=429, detail="Too many call requests for this number. Try again later.")
    call_id = store.create_direct_call(phone, req.name.strip(), req.employee_id, req.gender)
    result = await dial_employee({"id": call_id, "phone": phone, "name": req.name.strip(),
                                  "gender": req.gender, "employee_id": req.employee_id}, dial=dial)
    if result["status"] == "error":
        # Admin-only endpoint — show why, so a failed test call can be fixed without the server log.
        raise HTTPException(status_code=502, detail=f"Could not place the call — {result['detail']}")
    return {**result, "phone": phone, "call_id": call_id}


@app.post("/sgsp/call-status")
async def exotel_call_status(request: Request):
    form = await request.form()
    status, call_sid = (form.get("Status") or "").lower(), form.get("CallSid")
    log.info(f"[STATUS] CallSid={call_sid} Status={status} To={form.get('To')}")
    row = store.find_call_by_sid(call_sid) if call_sid else None
    if row and (status in UNANSWERED_STATUSES or (status == "completed" and row["status"] == "dialing")):
        info = pending_calls.get(row["phone"])
        if info and info.get("call_id") == row["id"]:
            await _discard_prewarm(pending_calls.pop(row["phone"], None))
        new_status = store.dial_failed(row["id"], status if status in UNANSWERED_STATUSES else "no_stream")
        if new_status:
            log.info(f"[CALL]  call {row['id']} not answered ({status}) → {new_status}")
    return {"status": "ok"}


@app.websocket("/sgsp/stream")
async def sgsp_phone_stream(ws: WebSocket):
    await ws.accept()
    log.info("[WS]    Exotel connection accepted")

    stream_sid = call_sid = phone = None
    try:
        async for raw in ws.iter_text():
            msg = json.loads(raw)
            if msg.get("event") == "start":
                sd = msg.get("start", {})
                stream_sid = sd.get("stream_sid") or sd.get("streamSid") or msg.get("stream_sid") or "unknown"
                call_sid = sd.get("call_sid") or sd.get("callSid") or "unknown"
                phone = sd.get("from") or sd.get("From") or "unknown"
                log.info(f"[WS]    start | call={call_sid} employee={phone} stream={stream_sid}")
                break
    except Exception as e:
        log.error(f"[WS]    Handshake error: {e}")
        await ws.close()
        return
    if phone is None:
        log.warning("[WS]    Stream closed before a start event")
        return

    normalized = normalize_phone(phone) or phone
    info = pending_calls.pop(normalized, None)
    inbound = info is None
    if inbound:
        # The employee dialled the ExoPhone back (e.g. after a missed call).
        employee = store.latest_employee_for_phone(normalized) or {}
        call_id = store.create_direct_call(normalized, employee.get("name"), employee.get("employee_id"),
                                           employee.get("gender"), status="in_call")
        # The identity step needs a name; an inbound caller is not asked to confirm it.
        employee = {"gender": employee.get("gender")}
        log.info(f"[WS]    Incoming call from {normalized} — answering as inbound (call {call_id})")
    else:
        call_id, employee = info["call_id"], info["employee"]
        store.mark_in_call(call_id)
        if call_sid == "unknown" and info.get("call_sid"):
            call_sid = info["call_sid"]

    conn = None
    if info and info.get("prewarm"):
        try:
            conn = await asyncio.wait_for(info["prewarm"], timeout=10)
        except Exception as e:
            log.warning(f"[LIVE]  Pre-warmed session unavailable: {e}")

    run_ts, call_dir = _new_call_dir(f"{normalized.lstrip('+')}_{call_id}")
    call_log = setup_call_logging(call_dir, run_ts)
    if conn:
        conn.log = call_log
    session = ExotelSession(ws, stream_sid, call_sid, call_log, run_ts, call_dir, normalized, employee, inbound=inbound)
    try:
        await session.run(conn)
    except Exception as e:
        call_log.error(f"[WS]    Session error: {e}", exc_info=True)
    finally:
        result = session.result()
        new_status = store.finish_call(call_id, result)
        log.info(f"[WS]    Call {call_id} ({normalized}) outcome={session.outcome} → {new_status}")
        call_log.info("[WS]    Connection closed")
        cleanup_call_logging(call_log)
        row = store.get_call(call_id) or {}
        await _post_webhook({**result, "callId": call_id, "campaignId": row.get("campaign_id"),
                             "employeeId": row.get("employee_id"), "attempt": row.get("attempts"),
                             "status": new_status})


# ── Campaigns ──────────────────────────────────────────────────────────────────
_COLUMN_ALIASES = {
    "phone": {"phone", "mobile", "mobile_no", "mobile_number", "phone_number", "phone_no", "contact", "contact_no"},
    "name": {"name", "employee_name", "emp_name", "full_name"},
    "employee_id": {"employee_id", "emp_id", "kgid", "kgid_no", "hrms_id", "employee_code"},
    "gender": {"gender", "sex"},
}


def _canonical(header) -> str | None:
    h = re.sub(r"[\s.\-/]+", "_", str(header or "").strip().lower()).strip("_")
    return next((key for key, names in _COLUMN_ALIASES.items() if h in names), None)


def parse_employee_file(file_name: str, content: bytes) -> tuple[list[dict], list[dict]]:
    """→ (valid unique rows, rejected rows with reasons). Needs phone + name columns."""
    if file_name.lower().endswith(".csv"):
        table = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        table = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    if not table:
        raise ValueError("The file is empty")
    columns = [_canonical(h) for h in table[0]]
    if "phone" not in columns or "name" not in columns:
        raise ValueError("The first row must have a phone/mobile column and a name column")

    rows, rejected, seen = [], [], set()
    for line_no, values in enumerate(table[1:], start=2):
        rec = {col: values[i] for i, col in enumerate(columns) if col and i < len(values)}
        if not any(v not in (None, "") for v in rec.values()):
            continue
        phone = normalize_phone(rec.get("phone"))
        name = str(rec.get("name") or "").strip()
        if not phone or not name:
            rejected.append({"row": line_no, "reason": "invalid phone" if not phone else "missing name"})
            continue
        if phone in seen:
            rejected.append({"row": line_no, "reason": "duplicate phone"})
            continue
        seen.add(phone)
        emp_id = rec.get("employee_id")
        rows.append({"phone": phone, "name": name,
                     "employee_id": str(int(emp_id)) if isinstance(emp_id, float) and emp_id.is_integer() else (str(emp_id) if emp_id else None),
                     "gender": str(rec.get("gender") or "").strip() or None})
    return rows, rejected


@app.post("/sgsp/campaigns", dependencies=[Depends(require_admin)])
async def create_campaign(name: str = Form(...), file: UploadFile = File(...), start_paused: bool = Form(False)):
    try:
        rows, rejected = parse_employee_file(file.filename or "", await file.read())
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not read the file: {e}")
    if not rows:
        raise HTTPException(status_code=422, detail=f"No valid rows ({len(rejected)} rejected)")
    campaign_id = store.create_campaign(name.strip() or file.filename, file.filename, rows, len(rejected))
    if start_paused:
        store.set_campaign_status(campaign_id, "paused")
    log.info(f"[CAMPAIGN] #{campaign_id} '{name}' — {len(rows)} employees, {len(rejected)} rejected")
    return {"campaign_id": campaign_id, "accepted": len(rows), "rejected": rejected[:200],
            "rejected_count": len(rejected)}


@app.get("/sgsp/campaigns", dependencies=[Depends(require_admin)])
async def list_campaigns():
    return {"campaigns": store.list_campaigns(), "active_calls": store.active_call_count(),
            "channel_limit": CHANNEL_LIMIT, "calling_hours_open": within_calling_hours()}


@app.get("/sgsp/campaigns/{campaign_id}", dependencies=[Depends(require_admin)])
async def get_campaign(campaign_id: int):
    stats = store.campaign_stats(campaign_id)
    if not stats:
        raise HTTPException(status_code=404, detail="No such campaign")
    return stats


@app.post("/sgsp/campaigns/{campaign_id}/{action}", dependencies=[Depends(require_admin)])
async def campaign_action(campaign_id: int, action: str):
    if action not in ("pause", "resume"):
        raise HTTPException(status_code=404, detail="Unknown action")
    if not store.set_campaign_status(campaign_id, "paused" if action == "pause" else "active"):
        raise HTTPException(status_code=404, detail="No such campaign")
    return store.campaign_stats(campaign_id)


_REPORT_COLUMNS = ["call_id", "employee_id", "name", "phone", "status", "attempts", "outcome", "outcome_source",
                   "enrollment_status", "bank_name", "objection", "employee_questions", "callback_time", "sms_sent",
                   "escalations", "duration_s", "cost_usd", "last_error", "call_dir", "updated_at"]


@app.get("/sgsp/campaigns/{campaign_id}/report.csv", dependencies=[Depends(require_admin)])
async def campaign_report(campaign_id: int):
    if not store.campaign_stats(campaign_id):
        raise HTTPException(status_code=404, detail="No such campaign")
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(_REPORT_COLUMNS)
    for c in store.campaign_calls(campaign_id):
        d = json.loads(c["details"] or "{}")
        esc = ", ".join(e.get("trigger") or "" for e in json.loads(c["escalations"] or "[]"))
        writer.writerow([
            c["id"], c["employee_id"], c["name"], c["phone"], c["status"], c["attempts"], c["outcome"],
            c["outcome_source"], c["enrollment_status"], d.get("bank_name"), d.get("objection") or d.get("notes"),
            d.get("employee_questions"), d.get("callback_time"), c["sms_sent"], esc, d.get("durationSeconds"),
            d.get("costUsd"), c["last_error"], c["call_dir"],
            datetime.fromtimestamp(c["updated_at"], IST).strftime("%Y-%m-%d %H:%M:%S"),
        ])
    # BOM so Excel opens the Kannada / UTF-8 text correctly.
    return StreamingResponse(iter(["﻿" + out.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="sgsp_campaign_{campaign_id}.csv"'})


class DncRequest(BaseModel):
    phone: str
    reason: str | None = None


@app.post("/sgsp/dnc", dependencies=[Depends(require_admin)])
async def add_dnc(req: DncRequest):
    phone = normalize_phone(req.phone)
    if not phone:
        raise HTTPException(status_code=422, detail="Invalid Indian mobile number")
    store.add_dnc(phone, req.reason or "added by admin")
    return {"status": "ok", "phone": phone}


# ── Pages ──────────────────────────────────────────────────────────────────────

@app.get("/sgsp/dashboard")
async def dashboard_page():
    return FileResponse(STATIC_DIR / "dashboard.html")


@app.get("/sgsp/web")
async def web_agent_page():
    return FileResponse(STATIC_DIR / "web_agent.html")


@app.websocket("/sgsp/web/stream")
async def sgsp_web_stream(ws: WebSocket):
    if ADMIN_TOKEN and ws.query_params.get("token") != ADMIN_TOKEN:
        await ws.close(code=4401)
        return
    await ws.accept()
    employee = {"name": (ws.query_params.get("name") or "").strip()[:80] or None,
                "gender": ws.query_params.get("gender") or None}
    log.info(f"[WEB]   Browser session started (employee name: {employee['name'] or 'none'})")
    run_ts, call_dir = _new_call_dir("web")
    call_log = setup_call_logging(call_dir, run_ts)
    session = WebSession(ws, call_log, run_ts, call_dir, employee=employee)
    try:
        await session.run()
        log.info(f"[WEB]   Web session {run_ts} outcome={session.outcome}")
    except Exception as e:
        call_log.error(f"[WEB]   Session error: {e}", exc_info=True)
    finally:
        call_log.info("[WEB]   Connection closed")
        cleanup_call_logging(call_log)
